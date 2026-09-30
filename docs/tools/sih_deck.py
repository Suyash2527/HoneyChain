"""Build the SIH 2026 idea presentation inside the official template (6 slides max, PDF for upload).
Same template, our content: template headings/logo/footer kept; detailed points, tables, diagrams, screenshots.
Usage: TEAM_NAME="..." TEAM_ID="..." python docs/tools/sih_deck.py <template.pptx> <shots_dir> <out.pptx>
"""
import os, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

TEMPLATE, SHOTS, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
TEAM = os.environ.get('TEAM_NAME', 'Your Team Name')
TEAM_ID = os.environ.get('TEAM_ID', '<Team ID>')

NAVY, BLUE = RGBColor(0x1F, 0x49, 0x7D), RGBColor(0x00, 0x70, 0xC0)
ORANGE, GREEN, RED = RGBColor(0xE9, 0x8B, 0x1D), RGBColor(0x2E, 0x8B, 0x57), RGBColor(0xB3, 0x37, 0x2F)
INK, GREY, LIGHT, LINE = RGBColor(0, 0, 0), RGBColor(0x59, 0x59, 0x59), RGBColor(0xF2, 0xF5, 0xFA), RGBColor(0xC9, 0xD4, 0xE6)
WHITE = RGBColor(255, 255, 255)
FONT = 'Arial'

prs = Presentation(TEMPLATE)

# ---------- helpers ----------
def drop_shape(shape):
    shape._element.getparent().remove(shape._element)

def add_runs(par, runs, size, color, bold=False, underline=False, italic=False):
    for r in runs:
        text, o = (r, {}) if isinstance(r, str) else r
        run = par.add_run(); run.text = text
        f = run.font; f.name = FONT; f.size = Pt(o.get('size', size)); f.bold = o.get('bold', bold)
        f.underline = o.get('underline', underline); f.italic = o.get('italic', italic); f.color.rgb = o.get('color', color)

def hang(par, indent, bullet=None, color=None):
    pPr = par._p.get_or_add_pPr()
    pPr.set('marL', str(int(Inches(indent)))); pPr.set('indent', str(-int(Inches(indent))))
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum', 'a:buClr'):
        for e in pPr.findall(qn(tag)): pPr.remove(e)
    if bullet:
        if color is not None:
            bc = etree.SubElement(pPr, qn('a:buClr')); c = etree.SubElement(bc, qn('a:srgbClr')); c.set('val', str(color))
        bu = etree.SubElement(pPr, qn('a:buChar')); bu.set('char', bullet)
    else:
        etree.SubElement(pPr, qn('a:buNone'))

def tb(slide, x, y, w, h, items, anchor=MSO_ANCHOR.TOP, fill=None, line=None, margin=0.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin); tf.margin_top = tf.margin_bottom = Inches(0.03)
    if fill is not None: box.fill.solid(); box.fill.fore_color.rgb = fill
    if line is not None: box.line.color.rgb = line; box.line.width = Pt(0.75)
    first = True
    for it in items:
        par = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        runs = it['runs'] if 'runs' in it else [it['text']]
        add_runs(par, runs, it.get('size', 10.5), it.get('color', INK), it.get('bold', False), it.get('underline', False), it.get('italic', False))
        par.alignment = it.get('align', PP_ALIGN.LEFT)
        par.space_after = Pt(it.get('after', 2)); par.space_before = Pt(it.get('before', 0)); par.line_spacing = it.get('ls', 1.0)
        if it.get('bullet'): hang(par, it.get('indent', 0.16), it['bullet'] if isinstance(it['bullet'], str) else '•', it.get('bcolor', BLUE))
        elif it.get('hang'): hang(par, it['hang'])
    return box

def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s

def label_shape(shape, text, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER):
    tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
    par = tf.paragraphs[0]; par.alignment = align
    add_runs(par, [text], size, color, bold)

def pic(slide, name, x, y, w=None, h=None, border=True):
    p = slide.shapes.add_picture(os.path.join(SHOTS, name), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if border: p.line.color.rgb = LINE; p.line.width = Pt(1)
    return p

def heading(slide, x, y, w, text, size=14, star=True):
    runs = ([('❖ ', {'color': NAVY})] if star else []) + [(text, {'color': NAVY, 'underline': True})]
    return tb(slide, x, y, w, 0.36, [{'runs': runs, 'size': size, 'bold': True}])

def sub(slide, x, y, w, text, size=12):
    return tb(slide, x, y, w, 0.3, [{'text': text, 'size': size, 'bold': True, 'color': NAVY, 'after': 0}])

def table(slide, x, y, w, cols, rows, size=9.5, head_size=10, row_h=0.3, first_bold=True):
    """Native table in the template's blue palette. cols = column widths (inches), rows[0] = header."""
    shape = slide.shapes.add_table(len(rows), len(cols), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    t = shape.table
    for i, cw in enumerate(cols): t.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        t.rows[r].height = Inches(row_h)
        for c, val in enumerate(row):
            cell = t.cell(r, c); cell.margin_left = cell.margin_right = Inches(0.05); cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame; tf.word_wrap = True
            par = tf.paragraphs[0]
            if r == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
                add_runs(par, [val], head_size, WHITE, True)
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT
                add_runs(par, [(val, {'bold': first_bold and c == 0, 'color': NAVY if (first_bold and c == 0) else INK})], size, INK)
    return shape

def find(slide, name):
    return next(s for s in slide.shapes if s.name == name)

def set_oval(slide):
    for s in slide.shapes:
        if s.name.startswith('Oval'):
            par = s.text_frame.paragraphs[0]
            for r in par.runs[1:]: r._r.getparent().remove(r._r)
            par.runs[0].text = TEAM; par.runs[0].font.size = Pt(9 if len(TEAM) > 14 else 11)

def set_title(slide, text, size):
    t = find(slide, 'Title 1'); par = t.text_frame.paragraphs[0]
    for r in par.runs: r.text = ''
    par.runs[-1].text = text; par.runs[-1].font.size = Pt(size)

def numbered(n, head, text):
    return {'runs': [(f'{n}. ', {'bold': True, 'color': NAVY}), (head + ' – ', {'bold': True}), text], 'size': 10, 'hang': 0.2, 'after': 2.5}

def bl(text, size=10, after=2, color=BLUE):
    return {'text': text, 'size': size, 'bullet': True, 'bcolor': color, 'after': after}

slides = list(prs.slides)

# =====================================================================
# 1. TITLE PAGE
# =====================================================================
s1 = slides[0]
box = find(s1, 'TextBox 9')
vals = {
    'Problem Statement ID': ('Problem Statement ID – ', 'SIH26021'),
    'Problem Statement Title': ('Problem Statement Title- ', 'Honey Chain: A block chain-based system for honey traceability and smart beekeeping management.'),
    'Theme': ('Theme- ', 'Agriculture, FoodTech & Rural Development'),
    'PS Category': ('PS Category- ', 'Software'),
    'Team ID': ('Team ID- ', TEAM_ID),
    'Team Name': ('Team Name (Registered on portal)- ', TEAM),
}
for par in box.text_frame.paragraphs:
    if not par.runs: continue
    key = next((k for k in vals if par.runs[0].text.startswith(k)), None)
    if not key: continue
    label, value = vals[key]
    for r in par.runs[1:]: r._r.getparent().remove(r._r)
    par.runs[0].text = label
    v = par.add_run(); v.text = value; v.font.bold = True; v.font.color.rgb = NAVY; v.font.name = par.runs[0].font.name
    par.space_after = Pt(5); par.line_spacing = 1.0; par.alignment = PP_ALIGN.LEFT
    par.runs[0].font.size = Pt(17); v.font.size = Pt(14 if key == 'Problem Statement Title' else 17)

# =====================================================================
# 2. IDEA TITLE - Proposed solution
# =====================================================================
s = slides[1]
set_title(s, 'HONEY CHAIN: Blockchain Honey Traceability + AI-IoT Smart Beekeeping', 19)
drop_shape(find(s, 'TextBox 8')); set_oval(s)
heading(s, 0.35, 1.25, 7.6, 'Proposed Solution (Describe your Idea/Solution/Prototype)')
tb(s, 0.35, 1.63, 7.6, 0.62, [{'runs': [('Honey Chain ', {'bold': True, 'color': NAVY}), 'gives every hive, harvest batch and jar a tamper-proof digital identity on a permissioned blockchain. Consumers verify authenticity by scanning a QR code; IoT + AI keep colonies healthy and productive.'], 'size': 10.5}])
sub(s, 0.35, 2.25, 7.6, '• Detailed explanation of the proposed solution')
tb(s, 0.35, 2.55, 7.6, 2.75, [
    numbered(1, 'Register', 'each KVIC bee box gets a Hive ID (owner, village, GPS, cluster, flora) written to the ledger.'),
    numbered(2, 'Sense', 'ESP32 hive node streams weight, brood temperature, humidity, colony sound, CO₂ and entrance activity; edge AI diagnoses varroa, foulbrood, queenless, pre-swarm and starvation and forecasts 30-day yield.'),
    numbered(3, 'Anchor', '14-day telemetry is hashed into a Merkle root and sealed on-chain; raw data stays local, so a 2G link is enough.'),
    numbered(4, 'Harvest', 'beekeeper picks hives, enters kg + flora → batch ID + printable QR; yield-per-hive plausibility is checked automatically.'),
    numbered(5, 'Certify', 'NABL lab records moisture, HMF, sucrose, reducing sugars, ash, C4 sugar, NMR profile and rice-syrup marker (SMR); pass/fail is computed from limits and cannot be edited.'),
    numbered(6, 'Move', 'processor / distributor / retailer handovers and retail packing are logged; every write prints a signed receipt.'),
    numbered(7, 'Verify', 'consumer scans the QR → trust score 0-100, origin, lab table, journey and validator signatures; forged codes return "no such batch".'),
])
sub(s, 0.35, 5.3, 7.6, '• Innovation and uniqueness of the solution')
tb(s, 0.35, 5.6, 7.6, 1.35, [
    bl('Lab-agnostic trust layer: anchors NMR, SMR, C4 and HMF results, so syrups engineered to pass basic tests are still rejected (CSE 2020)'),
    bl('Yield plausibility: claimed kg vs hive-weight telemetry exposes blending ("95 kg from 1 hive" is flagged)'),
    bl('Telemetry anchored by Merkle root: auditable claims with minimal bandwidth and cost'),
    bl('Permissioned PoA (KVIC / FSSAI / lab validators): no mining, no gas, accountable writers'),
    bl('Explainable AI with advice in English + Hindi; offline-first installable PWA built for villages'),
])
sub(s, 8.15, 1.25, 4.8, '• How it addresses the problem')
table(s, 8.15, 1.6, 4.8, [1.35, 3.45], [
    ['Problem (PS SIH26021)', 'Honey Chain answer'],
    ['Counterfeit honey', 'Lab results incl. NMR/SMR sealed per batch; failures are permanent; QR verifies each jar'],
    ['Low consumer trust', 'Scan-to-verify trust score, origin and journey read from the ledger, not from the seller'],
    ['Weak market linkage', 'Verified batches + custody trail support premium pricing and direct/retail buyers'],
    ['No traceability or hive-management support', 'Hive → harvest → jar trail; IoT-AI alerts, advice and yield forecast'],
], size=9.5, row_h=0.45)
sub(s, 8.15, 4.05, 4.8, '• Expected solution coverage (per problem statement)')
tb(s, 8.15, 4.35, 4.8, 1.3, [
    {'runs': [('✔ ', {'color': GREEN, 'bold': True}), ('Blockchain prototype + QR consumer authentication ', {'bold': True}), '– built and live'], 'size': 10, 'hang': 0.2, 'after': 3},
    {'runs': [('✔ ', {'color': GREEN, 'bold': True}), ('IoT hive monitoring + AI analytics ', {'bold': True}), '– built (simulated sensors, explainable model)'], 'size': 10, 'hang': 0.2, 'after': 3},
    {'runs': [('✔ ', {'color': GREEN, 'bold': True}), ('Scalable KVIC deployment framework ', {'bold': True}), '– phased plan, costs, risks (slide 4)'], 'size': 10, 'hang': 0.2, 'after': 3},
])
sub(s, 8.15, 5.5, 4.8, '• Trust score shown to the buyer (0-100)')
tb(s, 8.15, 5.8, 4.8, 1.1, [
    {'text': '+25 hives registered · +15 harvest logged · +40 lab PASS · +10 telemetry anchored · +10 custody / retail trail · −20 implausible yield per hive', 'size': 10, 'after': 3},
    {'text': 'Lab FAIL or score < 50 → REJECT.  Unknown / forged QR → "no such batch".  Ledger integrity is re-checked on every load.', 'size': 10, 'color': GREY},
], fill=LIGHT, margin=0.08)

# =====================================================================
# 3. TECHNICAL APPROACH
# =====================================================================
s = slides[2]
drop_shape(find(s, 'TextBox 8')); set_oval(s)
heading(s, 0.35, 1.25, 8.2, 'Technologies to be used (e.g. programming languages, frameworks, hardware)', size=13)
table(s, 0.35, 1.62, 8.3, [1.15, 4.35, 2.8], [
    ['Layer', 'Technologies', 'Role'],
    ['Frontend', 'React 18, Vite, JavaScript; installable offline-first PWA (service worker); html5-qrcode; EN / हिन्दी', 'Consumer, beekeeper, lab, supply-chain and officer portals'],
    ['Ledger', 'SHA-256 hash-chain (Web Crypto); Proof-of-Authority round-robin validators; Merkle trees', 'Tamper-evidence; integrity check on every load'],
    ['Smart contract', 'Solidity 0.8 (roles, batch, certify, custody) → Hyperledger Besu / Fabric-class network', 'Production ledger; a lab FAIL is final'],
    ['IoT hardware', 'ESP32, HX711 + load cell, DS18B20, DHT22, MEMS mic (I²S), CO₂ sensor, LoRa, solar + Li-ion', '15-min readings; 1 gateway per 20-40 hives'],
    ['AI / analytics', 'Feature model: brood-temp variance, sound shift, weight slope, activity drop, CO₂ → softmax; CNN (TFLite Micro)', '5 conditions + healthy; 30-day yield forecast (80% range)'],
    ['Data', 'On-chain: hashes only. Off-chain: time-series store; IPFS for lab PDFs', 'Privacy, low cost'],
    ['DevOps / security', 'GitHub Actions CI; Pages / Vercel / Docker + nginx; HTTPS; role-based keys (Ed25519 planned)', 'Repeatable, secure deployment'],
], size=9, head_size=9.5, row_h=0.36)
pic(s, 'desk-hive-detail.png', 8.85, 1.3, w=4.1)
tb(s, 8.85, 3.85, 4.1, 0.6, [{'runs': [('Working prototype – hive dashboard: ', {'bold': True, 'color': NAVY}), 'AI flags "Pre-swarm 98%" with advice; live sensor tiles and 30-day yield forecast.'], 'size': 9, 'color': GREY}])
heading(s, 0.35, 4.5, 12.6, 'Methodology and process for implementation (Flow Charts/Images/ working prototype)', size=13)
steps = [
    ('1 Register', ['Hive ID: owner, GPS, cluster, flora', 'Block HIVE_REGISTERED']),
    ('2 Sense', ['Node reads sensors every 15 min', 'Edge AI diagnosis + advice', 'Merkle root → SENSOR_ATTESTATION']),
    ('3 Harvest', ['Pick hives, kg, flora', 'Batch ID HC-YYYY-NNNN + QR', 'Yield/hive check (> 45 kg flagged)']),
    ('4 Certify', ['NABL lab enters 8 tests', 'Auto PASS / FAIL vs limits', 'LAB_CERTIFIED – immutable']),
    ('5 Move', ['CUSTODY_TRANSFER, RETAIL_PACKED', 'Who → whom, bottles, outlet', 'Signed receipt per write']),
    ('6 Verify', ['Scan QR or enter batch ID', 'Hash + signature re-check', 'Trust score, journey, verdict']),
]
for i, (h, pts) in enumerate(steps):
    x = 0.35 + i * 2.1
    sh = rect(s, x, 4.86, 2.0, 0.42, fill=NAVY if i < 5 else GREEN, shape=MSO_SHAPE.PENTAGON); label_shape(sh, h, 11.5)
    tb(s, x, 5.3, 2.0, 1.15, [{'text': p, 'size': 9, 'bullet': True, 'bcolor': BLUE, 'after': 1.5, 'indent': 0.12} for p in pts], fill=LIGHT, margin=0.05)
tb(s, 0.35, 6.5, 12.6, 0.42, [
    {'runs': [('Design rule: ', {'bold': True, 'color': NAVY}), 'only hashes and attestations go on-chain; raw sensor data stays local and is proven later via the Merkle root.  ', ('AI: ', {'bold': True, 'color': NAVY}), 'softmax over 6 states, 24 / 24 injected conditions identified in tests.  ', ('Data: ', {'bold': True, 'color': NAVY}), 'synthetic records modelled on real districts, flora and public parameters (slide 6).'], 'size': 9, 'color': INK}], fill=None)

# =====================================================================
# 4. FEASIBILITY AND VIABILITY
# =====================================================================
s = slides[3]
drop_shape(find(s, 'TextBox 8')); set_oval(s)
cols = [
    ('Analysis of the feasibility of the idea', GREEN, [
        ('Technical: ', 'working prototype live (ledger, QR verify, lab portal with NMR/SMR, hive AI, receipts, PWA); mature open-source stack'),
        ('Hardware: ', 'ESP32 node ≈ ₹1.5-2.5k per hive*; one LoRa gateway serves 20-40 hives within ~5 km'),
        ('Economic: ', 'near-zero ledger cost (PoA, no gas); QR label ≈ ₹0.1-0.3 per jar*'),
        ('Operational: ', 'fits NBHM / KVIC clusters, NABL labs and FSSAI oversight; no new institution needed'),
        ('Regulatory: ', 'minimal personal data (name, village); limits FSSAI/BIS-style, final limits set by the notified lab'),
        ('Scalable: ', 'validator nodes per KVIC region; hashes on-chain, raw data off-chain')]),
    ('Potential challenges and risks', RED, [
        ('Garbage-in: ', 'a dishonest lab or beekeeper can enter false data'),
        ('Adulteration that beats basic tests: ', 'engineered syrups pass C4/HMF (CSE 2020)'),
        ('Connectivity and literacy: ', 'patchy 2G/4G, low digital skills'),
        ('QR cloning: ', 'copying a genuine label onto fake jars, refilling'),
        ('Field hardware: ', 'sensor drift, theft, power, weather'),
        ('Adoption and governance: ', 'cluster buy-in, cost sharing, who runs validators and keys')]),
    ('Strategies for overcoming these challenges', NAVY, [
        ('Trust in inputs: ', 'accredited-lab keys, random re-tests, IoT weight cross-check vs claimed yield'),
        ('Lab-agnostic: ', 'anchor NMR + SMR alongside C4/HMF; new tests plug in'),
        ('Village-ready: ', 'offline-first PWA, EN/HI, QR-first flows, assisted entry via KVIC field officers'),
        ('Anti-cloning: ', 'per-jar serials + single-use claim codes, tamper seals (roadmap)'),
        ('Hardware: ', 'calibration, solar, IP-rated enclosure, sensor-anomaly flags'),
        ('Governance: ', 'multi-stakeholder validators (KVIC, FSSAI, labs); hardware key storage; pilot via SHG trainers')]),
]
for i, (h, c, pts) in enumerate(cols):
    x = 0.35 + i * 4.24
    hd = rect(s, x, 1.28, 4.1, 0.4, fill=c); label_shape(hd, h, 11.5)
    tb(s, x, 1.68, 4.1, 3.0, [{'runs': [(a, {'bold': True, 'color': c if c != NAVY else NAVY}), b], 'size': 10.3, 'bullet': True, 'bcolor': c, 'after': 4, 'indent': 0.14} for a, b in pts], fill=LIGHT, margin=0.08)
sub(s, 0.35, 4.75, 6.2, 'Unit economics – planning estimates* (to be validated in the pilot)')
table(s, 0.35, 5.05, 6.2, [2.6, 1.7, 1.9], [
    ['Item', 'Estimate', 'Note'],
    ['Hive node (ESP32 + sensors)', '₹1.5-2.5k / hive', 'solar + LoRa'],
    ['LoRa gateway (per cluster)', '₹15-25k', '20-40 hives each'],
    ['QR label per jar', '₹0.1-0.3', 'printed by beekeeper'],
    ['Pilot: 100 hives, 3-5 gateways', '≈ ₹2-3.8 lakh', 'hardware only'],
], size=9, head_size=9.5, row_h=0.31)
sub(s, 6.75, 4.75, 6.2, 'Phased rollout and validation plan')
ph = [('Pilot', '0-3 mo', '2 clusters · 50 beekeepers · 100 hives · 1 lab'), ('State', '4-9 mo', '10 clusters · 3 states · QR at Khadi outlets'), ('Regional', '10-18 mo', 'all Honey Mission states · FSSAI dashboard'), ('National', 'Yr 2+', 'open API · export certificates · trained CNN')]
for i, (n, t, d) in enumerate(ph):
    x = 6.75 + i * 1.56
    ch = rect(s, x, 5.05, 1.5, 0.4, fill=[BLUE, RGBColor(0x2F, 0x6B, 0x92), GREEN, ORANGE][i], shape=MSO_SHAPE.PENTAGON); label_shape(ch, f'{n} · {t}', 9, True, WHITE if i < 3 else INK)
    tb(s, x, 5.47, 1.5, 0.62, [{'text': d, 'size': 8.5, 'color': INK, 'align': PP_ALIGN.CENTER}])
tb(s, 6.75, 6.12, 6.2, 0.8, [
    bl('Validation first: 10-minute interviews with 5 beekeepers, 2 NABL labs, 1 KVIC field officer (script prepared)', 9.5, 2),
    bl('Pilot KPIs: % batches with QR + lab certificate, scans/month, price realisation, colony-loss rate', 9.5, 2),
    {'text': '*Planning estimates, not quotes.', 'size': 8.5, 'color': GREY},
], fill=LIGHT, margin=0.08)

# =====================================================================
# 5. IMPACT AND BENEFITS
# =====================================================================
s = slides[4]
drop_shape(find(s, 'TextBox 8')); set_oval(s)
heading(s, 0.35, 1.25, 8.4, 'Potential impact on the target audience', size=13)
table(s, 0.35, 1.62, 8.4, [1.5, 3.0, 3.9], [
    ['Audience', 'Today', 'With Honey Chain'],
    ['Rural beekeepers (KVIC Honey Mission)', 'Cannot prove purity; sell to middlemen at commodity prices; late disease detection', 'Verifiable quality → direct buyers and premium price; early alerts and yield forecast reduce colony loss'],
    ['Consumers', 'Cannot tell real honey from syrup blends; labels are easy to copy', 'Scan-to-verify in seconds: origin, lab report, journey, trust score; fake batches flagged before purchase'],
    ['KVIC / FSSAI / regulators', 'No live view of production or adulteration; manual audits', 'Cluster dashboards, tamper-evident audit trail, adulteration flags, hotspot targeting for extension'],
    ['Processors, retailers, e-commerce', 'Brand and recall risk from unverifiable supply', 'Trusted supply chain with custody trail; verified, premium-ready listings'],
    ['Labs and exporters', 'Reports on paper, easy to dispute or re-use', 'Signed, immutable certificates (NMR/SMR/C4/HMF) reusable as export evidence'],
], size=9, head_size=9.5, row_h=0.5)
heading(s, 0.35, 4.72, 8.4, 'Benefits of the solution (social, economic, environmental, etc.)', size=13)
ben = [('Social', GREEN, ['Trust in rural produce', 'Livelihood security incl. SHGs / women beekeepers', 'Bilingual, low-literacy design']),
       ('Economic', ORANGE, ['Higher realisation per kg', 'Less fraud loss for buyers', 'Export-ready traceability', 'Lower colony-loss cost']),
       ('Environmental', BLUE, ['Healthier colonies → pollination', 'Targeted treatment, less chemical', 'Low-energy ledger (no mining)', 'Paperless certificates']),
       ('Governance', NAVY, ['Tamper-evident audit trail', 'Data-driven KVIC planning', 'Aligned to SDG 1, 2, 8, 12, 15'])]
for i, (h, c, pts) in enumerate(ben):
    x = 0.35 + i * 2.13
    hd = rect(s, x, 5.08, 2.05, 0.32, fill=c); label_shape(hd, h, 10.5, True, WHITE if c != ORANGE else INK)
    tb(s, x, 5.4, 2.05, 1.5, [bl(p, 9.5, 2.5, c) for p in pts], fill=LIGHT, margin=0.06)
sub(s, 9.0, 1.25, 4.0, 'WHY NOW – public evidence', 12)
call = [('17 of 22', 'honey samples tested were adulterated with sugar syrup (77%)', 'CSE investigation, Dec 2020', RED),
        ('46%', 'of imported honey in the EU was suspected of adulteration in 2022 (EU coordinated action)', 'via Frontiers Sust. Food Syst., 2025', RED),
        ('1.4 lakh MT', 'natural honey produced in India in 2024; ~1.07 lakh MT exported in FY 2023-24', 'PIB - National Beekeeping & Honey Mission', GREEN),
        ('₹500 crore', 'NBHM outlay, FY 2020-21 to 2025-26: production is growing, trust must keep pace', 'PIB', BLUE)]
for i, (n, d, src, c) in enumerate(call):
    y = 1.6 + i * 1.02
    rect(s, 9.0, y, 0.07, 0.92, fill=c)
    tb(s, 9.15, y - 0.03, 3.85, 1.0, [{'text': n, 'size': 19, 'bold': True, 'color': c, 'after': 0}, {'text': d, 'size': 9.2, 'after': 1}, {'text': src, 'size': 8, 'color': GREY}])
sub(s, 9.0, 5.72, 4.0, 'Pilot KPIs (baseline measured at pilot start)', 11)
tb(s, 9.0, 6.0, 3.95, 0.92, [
    bl('100% of pilot batches with QR + lab certificate', 9, 1.5),
    bl('Scans / month; adulterated batches flagged per 1,000', 9, 1.5),
    bl('Price realisation vs unlabelled honey; colony-loss rate', 9, 1.5),
], fill=LIGHT, margin=0.06)

# =====================================================================
# 6. RESEARCH AND REFERENCES
# =====================================================================
s = slides[5]
drop_shape(find(s, 'TextBox 8')); set_oval(s)
heading(s, 0.35, 1.25, 12.6, 'Details / Links of the reference and research work', size=13)

def ref(n, title, url, color=NAVY):
    return [{'runs': [(f'{n}. ', {'bold': True, 'color': color}), (title, {'bold': True})], 'size': 9.5, 'hang': 0.22, 'after': 0},
            {'text': url, 'size': 8, 'color': BLUE, 'hang': 0.22, 'after': 3.5}]

left = []
left += [{'text': 'Market, regulatory and standards evidence', 'size': 11, 'bold': True, 'color': NAVY, 'after': 3}]
left += ref(1, 'CSE – "Laboratory results of honey testing" (Dec 2020): NMR tests on branded honey', 'cdn.downtoearth.org.in/report_laboratory_results_of_honey_testing_20201201.pdf')
left += ref(2, 'FSSAI press release reviewing the CSE honey-sample issue (3 Dec 2020)', 'fssai.gov.in/upload/press_release/2020/12/5fc8eb6d70965Press_Release_Honey_Sample_03_12_2020.pdf')
left += ref(3, 'PIB – National Beekeeping & Honey Mission: outlay, production, exports', 'pib.gov.in/PressReleasePage.aspx?PRID=2185400   |   National Bee Board: nbb.gov.in')
left += ref(4, 'The Tribune – "Honeytrap": syrups designed to bypass purity tests', 'tribuneindia.com/news/nation/honeytrap-cse-investigations-claim-nefarious-adulteration-business-of-honey-designed-to-bypass-purity-tests-178930')
left += ref(5, 'Business Standard – FSSAI on the SMR (rice-syrup marker) test', 'business-standard.com/article/companies/fssai-questions-cse-for-not-conducting-smr-test-on-honey-brands-120120301199_1.html')
left += ref(6, 'Codex Alimentarius – Standard for Honey, CXS 12-1981 (rev. 2001; amended 2019, 2022)', 'fao.org/input/download/standards/310/cxs_012e.pdf   |   Indian standard: BIS IS 4941 (honey)')
tb(s, 0.35, 1.6, 6.3, 3.85, left)
right = []
right += [{'text': 'Technology and research', 'size': 11, 'bold': True, 'color': NAVY, 'after': 3}]
right += ref(7, 'Ramsey et al. – "The prediction of swarming in honeybee colonies using vibrational spectra", Sci. Rep. 10, 9798 (2020)', 'nature.com/articles/s41598-020-66115-5')
right += ref(8, 'Kulyukin – "Audio, image, video, and weather datasets for continuous electronic beehive monitoring", Applied Sciences 11(10), 4632 (2021)', 'Open datasets for training the production CNN')
right += ref(9, 'Machine Learning and Computer Vision Techniques in Continuous Beehive Monitoring Applications: A survey', 'arxiv.org/pdf/2208.00085')
right += ref(10, '"Fujairah Honey Chain (FHC): A Blockchain Framework for Monitoring Honey Production", Information 16(8), 626 (2025)', 'mdpi.com/2078-2489/16/8/626')
right += ref(11, '"Analysing blockchain adoption in beekeeping: application of theoretical models", Frontiers in Sustainable Food Systems (2025)', 'frontiersin.org/journals/sustainable-food-systems/articles/10.3389/fsufs.2025.1566341/full')
tb(s, 6.85, 1.6, 6.1, 3.85, right)
sub(s, 0.35, 4.42, 12.6, 'Data used in the prototype (transparency)', 11)
tb(s, 0.35, 4.68, 12.6, 0.9, [
    {'runs': [('Real public data: ', {'bold': True, 'color': NAVY}), 'CSE 2020 test findings, PIB / NBHM production and outlay figures, FSSAI / BIS / Codex honey parameters (refs 1-6) - used for the problem evidence, lab limits and impact statistics.'], 'size': 9.5, 'bullet': True, 'bcolor': GREEN, 'after': 2},
    {'runs': [('Prototype dataset: ', {'bold': True, 'color': NAVY}), '15 hives, 10 batches and 51 ledger blocks across 8 real districts in 6 states (Bharatpur, Muzaffarpur, Ludhiana, South 24 Parganas, Saharanpur, Nashik, Nagpur, Sawai Madhopur) using real flowering regions. Beekeeper names, hive IDs, lab reports and sensor streams are synthetic (no real individuals); real pilot data replaces them in Phase 0.'], 'size': 9.5, 'bullet': True, 'bcolor': ORANGE, 'after': 2},
    {'runs': [('AI data path: ', {'bold': True, 'color': NAVY}), 'prototype model runs on simulated telemetry; production CNN to be trained on open beehive datasets (refs 8, 9) plus KVIC apiary recordings.'], 'size': 9.5, 'bullet': True, 'bcolor': BLUE},
], fill=LIGHT, margin=0.08)
sub(s, 0.35, 5.6, 7.4, 'Research takeaways applied in the design', 11)
tb(s, 0.35, 5.88, 7.5, 1.02, [
    bl('Engineered syrups can pass C4/HMF but fail NMR → ledger is lab-agnostic and anchors NMR + SMR results (refs 1, 4, 5)', 8.6, 1),
    bl('Hive vibration/acoustics can forewarn swarming (ref 7) → colony-sound feature in the model; CNN trained on open datasets (refs 8, 9) is the production path', 8.6, 1),
    bl('Blockchain honey frameworks exist (refs 10, 11); Honey Chain focuses on NMR/SMR-aware labs, Merkle-anchored telemetry and an offline bilingual PWA for Indian clusters', 8.6, 1),
    bl('Prototype lab limits are indicative (FSSAI / BIS / Codex style); final limits come from the notified lab', 8.6, 1),
], fill=LIGHT, margin=0.08)
rect(s, 8.05, 5.62, 4.9, 1.3, fill=None, line=LINE)
tb(s, 8.1, 5.64, 4.8, 0.24, [{'text': 'TRY THE PROTOTYPE', 'size': 10, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}])
pic(s, 'qr-demo.png', 8.35, 5.9, w=0.76, border=False); pic(s, 'qr-repo.png', 9.95, 5.9, w=0.76, border=False); pic(s, 'qr-verify.png', 11.55, 5.9, w=0.76, border=False)
tb(s, 8.05, 6.7, 1.6, 0.22, [{'text': 'Live demo', 'size': 8.5, 'bold': True, 'align': PP_ALIGN.CENTER}])
tb(s, 9.65, 6.7, 1.6, 0.22, [{'text': 'Source code', 'size': 8.5, 'bold': True, 'align': PP_ALIGN.CENTER}])
tb(s, 11.25, 6.7, 1.6, 0.22, [{'text': 'Verify sample', 'size': 8.5, 'bold': True, 'align': PP_ALIGN.CENTER}])
tb(s, 1.35, 7.0, 0.1, 0.1, [{'text': '', 'size': 1}])

# drop the "Important instructions" slide
sldIdLst = prs.slides._sldIdLst
last = sldIdLst[-1]
prs.part.drop_rel(last.rId); sldIdLst.remove(last)
prs.save(OUT)
print('saved', OUT, 'slides:', len(prs.slides))
